"""Minimal tracked-change helpers on top of python-docx (lxml)."""
import copy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches

AUTHOR = "Claude"
DATE = "2026-10-14T12:00:00Z"
_next = [90000]
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def nid():
    _next[0] += 1
    return str(_next[0])


def _rev(tag):
    e = OxmlElement(tag)
    e.set(qn("w:id"), nid())
    e.set(qn("w:author"), AUTHOR)
    e.set(qn("w:date"), DATE)
    return e


def ptext(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def make_run(text, italic=False, bold=False, sup=False, sub=False, rpr_from=None):
    r = OxmlElement("w:r")
    if rpr_from is not None and rpr_from.find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(rpr_from.find(qn("w:rPr")))
        for tag in ("w:i", "w:iCs", "w:b", "w:bCs", "w:vertAlign", "w:rStyle"):
            for x in rpr.findall(qn(tag)):
                rpr.remove(x)
    else:
        rpr = OxmlElement("w:rPr")
    # Schema order inside rPr: rStyle, rFonts, b, bCs, i, iCs, ..., vertAlign ...
    pos = 0
    if rpr.find(qn("w:rFonts")) is not None:
        pos = list(rpr).index(rpr.find(qn("w:rFonts"))) + 1
    adds = []
    if bold:
        adds += ["w:b", "w:bCs"]
    if italic:
        adds += ["w:i", "w:iCs"]
    for k, tag in enumerate(adds):
        rpr.insert(pos + k, OxmlElement(tag))
    if sup or sub:
        va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript" if sup else "subscript")
        # vertAlign comes after sz/szCs/highlight/u/effect etc.; append is valid when those precede it
        rpr.append(va)
    if len(rpr):
        r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(XML_SPACE, "preserve")
    r.append(t)
    return r


def runs_from_spec(spec, rpr_from=None):
    """spec: list of str or (str, {"italic":..,"sup":..,"bold":..})."""
    out = []
    for s in spec:
        if isinstance(s, str):
            out.append(make_run(s, rpr_from=rpr_from))
        else:
            out.append(make_run(s[0], rpr_from=rpr_from, **s[1]))
    return out


def ins_wrap(runs):
    ins = _rev("w:ins")
    for r in runs:
        ins.append(r)
    return ins


def del_run(r):
    for t in r.findall(qn("w:t")):
        t.tag = qn("w:delText")
        t.set(XML_SPACE, "preserve")
    d = _rev("w:del")
    r.addprevious(d)
    d.append(r)
    return d


def _isolate(p):
    """Split runs (direct children of p) so each content child sits in its own run."""
    for r in list(p.findall(qn("w:r"))):
        content = [c for c in r if c.tag != qn("w:rPr")]
        if len(content) <= 1:
            continue
        rpr = r.find(qn("w:rPr"))
        for c in content:
            nr = OxmlElement("w:r")
            if rpr is not None:
                nr.append(copy.deepcopy(rpr))
            nr.append(c)
            r.addprevious(nr)
        p.remove(r)


def _split_run(r, offset):
    """Split a single-w:t run at char offset; returns (left, right)."""
    t = r.find(qn("w:t"))
    txt = t.text or ""
    right = copy.deepcopy(r)
    t.text = txt[:offset]; t.set(XML_SPACE, "preserve")
    rt = right.find(qn("w:t")); rt.text = txt[offset:]; rt.set(XML_SPACE, "preserve")
    r.addnext(right)
    return r, right


def _text_runs(p):
    return [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]


def tracked_replace(p, old, new_spec, rpr_from=None):
    """Track-delete the exact text `old` (direct runs of p) and insert new_spec after it.
    If old == "" , new_spec is appended at the end of p (tracked insertion)."""
    _isolate(p)
    if old == "":
        runs = runs_from_spec(new_spec, rpr_from=rpr_from if rpr_from is not None else (_text_runs(p) or [None])[-1])
        p.append(ins_wrap(runs))
        return
    full = "".join(r.find(qn("w:t")).text or "" for r in _text_runs(p))
    start = full.find(old)
    if start < 0:
        raise ValueError("text not found in paragraph: %r" % old[:80])
    end = start + len(old)
    # split at boundaries
    for boundary in (start, end):
        pos = 0
        for r in _text_runs(p):
            L = len(r.find(qn("w:t")).text or "")
            if pos < boundary < pos + L:
                _split_run(r, boundary - pos)
                break
            pos += L
    pos = 0
    victims = []
    for r in _text_runs(p):
        L = len(r.find(qn("w:t")).text or "")
        if pos >= start and pos + L <= end and L > 0:
            victims.append(r)
        pos += L
    tmpl = rpr_from if rpr_from is not None else victims[0]
    last = None
    for r in victims:
        last = del_run(r)
    if new_spec:
        last.addnext(ins_wrap(runs_from_spec(new_spec, rpr_from=tmpl)))


def tracked_insert_after(p, old_anchor, new_spec, rpr_from=None):
    """Insert new_spec immediately after the exact text old_anchor (no deletion)."""
    _isolate(p)
    full = "".join(r.find(qn("w:t")).text or "" for r in _text_runs(p))
    k = full.find(old_anchor)
    if k < 0:
        raise ValueError("anchor not found: %r" % old_anchor[:80])
    end = k + len(old_anchor)
    pos = 0
    for r in _text_runs(p):
        L = len(r.find(qn("w:t")).text or "")
        if pos < end < pos + L:
            _split_run(r, end - pos)
            break
        pos += L
    pos = 0
    for r in _text_runs(p):
        L = len(r.find(qn("w:t")).text or "")
        pos += L
        if pos == end:
            r.addnext(ins_wrap(runs_from_spec(new_spec, rpr_from=rpr_from if rpr_from is not None else r)))
            return
    raise ValueError("could not place insertion")


def new_paragraph(ppr_from, spec, rpr_from=None, keep_style=True):
    """New tracked-inserted paragraph; pPr copied from ppr_from (paragraph element)."""
    p = OxmlElement("w:p")
    src = ppr_from.find(qn("w:pPr")) if ppr_from is not None else None
    ppr = copy.deepcopy(src) if src is not None else OxmlElement("w:pPr")
    for x in ppr.findall(qn("w:rPr")):
        ppr.remove(x)
    for x in ppr.findall(qn("w:sectPr")) + ppr.findall(qn("w:pPrChange")):
        ppr.remove(x)
    if not keep_style:
        for x in ppr.findall(qn("w:pStyle")):
            ppr.remove(x)
    prpr = OxmlElement("w:rPr"); prpr.append(_rev("w:ins"))
    ppr.append(prpr)
    p.append(ppr)
    if spec:
        p.append(ins_wrap(runs_from_spec(spec, rpr_from=rpr_from)))
    return p


def picture_paragraph(doc, ppr_from, image_path, width_in):
    """Tracked-inserted paragraph holding an inline picture (relationship created by python-docx)."""
    tmp = doc.add_paragraph()
    run = tmp.add_run()
    run.add_picture(image_path, width=Inches(width_in))
    r = run._r
    tmp._p.getparent().remove(tmp._p)
    p = new_paragraph(ppr_from, None)
    p.append(ins_wrap([r]))
    return p


def table_from_rows(doc, rows, header_bold=True, font_sz=16):
    """Tracked-inserted table (rows: list of lists of str). Returns the w:tbl element."""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = None
    tbl = t._tbl
    tbl.getparent().remove(tbl)
    tblPr = tbl.find(qn("w:tblPr"))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "insideH"):
        b = OxmlElement("w:" + side)
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "auto")
        borders.append(b)
    # Schema order: tblStyle, ..., tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
    after = None
    for tag in ("w:tblInd", "w:tblCellSpacing", "w:jc", "w:tblW", "w:tblStyle"):
        after = tblPr.find(qn(tag))
        if after is not None:
            break
    if after is not None:
        after.addnext(borders)
    else:
        tblPr.insert(0, borders)
    for i, (tr, vals) in enumerate(zip(tbl.findall(qn("w:tr")), rows)):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr"); tr.insert(0, trpr)
        trpr.append(_rev("w:ins"))
        for tc, val in zip(tr.findall(qn("w:tc")), vals):
            p = tc.find(qn("w:p"))
            ppr = OxmlElement("w:pPr")
            prpr = OxmlElement("w:rPr"); prpr.append(_rev("w:ins")); ppr.append(prpr)
            p.insert(0, ppr)
            r = make_run(str(val), bold=(header_bold and i == 0))
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(font_sz))
            szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(font_sz))
            rpr.append(sz); rpr.append(szcs)
            p.append(ins_wrap([r]))
    return tbl


def comment_markers(first_el, last_el, cid):
    """Wrap commentRangeStart before first_el and End + reference after last_el (siblings in a p)."""
    s = OxmlElement("w:commentRangeStart"); s.set(qn("w:id"), str(cid))
    e = OxmlElement("w:commentRangeEnd"); e.set(qn("w:id"), str(cid))
    ref_r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr"); rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "CommentReference"); rpr.append(rs)
    ref_r.append(rpr)
    cr = OxmlElement("w:commentReference"); cr.set(qn("w:id"), str(cid)); ref_r.append(cr)
    first_el.addprevious(s)
    last_el.addnext(e)
    e.addnext(ref_r)
